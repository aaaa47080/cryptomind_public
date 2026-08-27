#!/usr/bin/env python3
"""將 Manifund 計劃書 markdown 同步產生為 DOCX。

來源：docs/plans/2026-08-14-general-agent-platform-ai-studio-manifund-design.md
輸出：docs/plans/2026-08-14-general-agent-platform-ai-studio-manifund-plan.docx

設計：
- 封面頁（kicker + 標題 + 中繼資料表）→ 分頁 → 目錄（TOC 欄位，Word 開啟後更新）→ 分頁 → 正文
- §6 架構圖重用原圖 docs/plans/assets/manifund-architecture.png（若存在）
- 章節標題用 Heading 1/2/3 樣式，程式碼區塊用等寬灰底，表格用 Table Grid + 表頭底色
- 僅讀取 markdown，不修改任何產品代碼或 schema
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs/plans/2026-08-14-general-agent-platform-ai-studio-manifund-design.md"
DST = ROOT / "docs/plans/2026-08-14-general-agent-platform-ai-studio-manifund-plan.docx"
ARCH_IMG = ROOT / "docs/plans/assets/manifund-architecture.png"

LATIN = "Calibri"
CJK = "Microsoft JhengHei"
MONO = "Consolas"
ACCENT_H1 = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT_H2 = RGBColor(0x2E, 0x5A, 0x88)
ACCENT_H3 = RGBColor(0x3A, 0x3A, 0x3A)
LINK = RGBColor(0x1A, 0x56, 0xC4)
GREY = RGBColor(0x80, 0x80, 0x80)


# ----------------------------- markdown 解析 -----------------------------
def parse_md(text: str):
    lines = text.split("\n")
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.rstrip()

        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            buf = []
            i += 1
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            blocks.append({"type": "code", "lang": lang, "lines": buf})
            continue

        if re.match(r"^#{1,6} ", stripped):
            m = re.match(r"^(#{1,6}) (.+)$", stripped)
            level = len(m.group(1))
            blocks.append({"type": f"h{level}", "text": m.group(2).strip()})
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|?\s*$", lines[i + 1]):
            rows = []
            while i < n and lines[i].startswith("|"):
                cell_line = lines[i].strip()
                if re.match(r"^\|[\s:|-]+\|?\s*$", cell_line):
                    i += 1
                    continue  # separator
                cells = [c.strip() for c in cell_line.strip("|").split("|")]
                rows.append(cells)
                i += 1
            blocks.append({"type": "table", "rows": rows})
            continue

        if stripped in ("---", "***", "___"):
            blocks.append({"type": "hr"})
            i += 1
            continue

        m_num = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_num:
            items = []
            while i < n:
                m2 = re.match(r"^(\d+)\.\s+(.*)$", lines[i].rstrip())
                if m2:
                    items.append((int(m2.group(1)), m2.group(2)))
                    i += 1
                    continue
                break
            blocks.append({"type": "numbered", "items": items})
            continue

        m_bul = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m_bul:
            items = []
            while i < n:
                mb = re.match(r"^(\s*)[-*]\s+(.*)$", lines[i])
                if mb:
                    lvl = 0 if len(mb.group(1)) < 2 else 1
                    items.append((lvl, mb.group(2).rstrip()))
                    i += 1
                    continue
                break
            blocks.append({"type": "bullets", "items": items})
            continue

        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].lstrip().startswith(">"):
                buf.append(lines[i].lstrip()[1:].strip())
                i += 1
            blocks.append({"type": "quote", "text": " ".join(x for x in buf if x)})
            continue

        if stripped == "":
            i += 1
            continue

        # ordinary paragraph (may span consecutive non-empty lines)
        buf = [stripped]
        i += 1
        while i < n and lines[i].strip() and not _is_block_start(lines[i], lines[i + 1] if i + 1 < n else ""):
            buf.append(lines[i].strip())
            i += 1
        blocks.append({"type": "para", "text": " ".join(buf)})
    return blocks


def _is_block_start(cur: str, nxt: str) -> bool:
    s = cur.rstrip()
    if re.match(r"^#{1,6} ", s):
        return True
    if s.startswith("```"):
        return True
    if s.startswith("|") and re.match(r"^\|[\s:|-]+\|?\s*$", nxt):
        return True
    if s in ("---", "***", "___"):
        return True
    if re.match(r"^(\d+)\.\s+", s):
        return True
    if re.match(r"^(\s*)[-*]\s+", cur):
        return True
    if s.lstrip().startswith(">"):
        return True
    return False


# ----------------------------- 文件輔助 -----------------------------
def set_cjk_font(run, latin=LATIN, cjk=CJK, mono=False):
    if mono:
        latin = MONO
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)


def add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    # CT_RPr schema 順序：rFonts → color → u
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), LATIN)
    rfonts.set(qn("w:hAnsi"), LATIN)
    rfonts.set(qn("w:eastAsia"), CJK)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A56C4")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


# 行內格式：**bold**、`code`、[text](url)、裸 URL
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s）)、。；<>]+)")


def add_inline(paragraph, text: str, base_size=10.5, mono=False):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_cjk_font(r, mono=mono)
            r.font.size = Pt(base_size)
        tok = m.group(0)
        if tok.startswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
            set_cjk_font(r, mono=mono)
            r.font.size = Pt(base_size)
        elif tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            set_cjk_font(r, mono=True)
            r.font.size = Pt(base_size)
        elif tok.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            add_hyperlink(paragraph, mm.group(2), mm.group(1))
        elif tok.startswith("http"):
            add_hyperlink(paragraph, tok, tok)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_cjk_font(r, mono=mono)
        r.font.size = Pt(base_size)


def shade_cell(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def para_shade(paragraph, hex_color):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    ppr.append(shd)


def add_toc(doc):
    p = doc.add_paragraph()
    r1 = p.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    r1._r.append(f1)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-1" \\h \\z \\u'
    r1._r.append(it)
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    r1._r.append(f2)
    note = doc.add_paragraph()
    rn = note.add_run("（請以 Word 開啟後，在目錄上按右鍵 → 更新功能變數，以顯示頁碼）")
    rn.italic = True
    rn.font.size = Pt(9)
    rn.font.color.rgb = GREY
    set_cjk_font(rn)
    p2 = doc.add_paragraph()
    r2 = p2.add_run()
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    r2._r.append(f3)


# ----------------------------- 組裝 -----------------------------
def build():
    text = SRC.read_text(encoding="utf-8")
    blocks = parse_md(text)

    doc = Document()
    # 預設樣式
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), LATIN)
    rfonts.set(qn("w:hAnsi"), LATIN)
    rfonts.set(qn("w:eastAsia"), CJK)

    for hname, sz, color in (("Heading 1", 18, ACCENT_H1), ("Heading 2", 14, ACCENT_H2), ("Heading 3", 12, ACCENT_H3)):
        st = doc.styles[hname]
        st.font.size = Pt(sz)
        st.font.color.rgb = color
        st.font.bold = True
        srpr = st.element.get_or_add_rPr()
        srfonts = srpr.find(qn("w:rFonts"))
        if srfonts is None:
            srfonts = OxmlElement("w:rFonts")
            srpr.append(srfonts)
        srfonts.set(qn("w:ascii"), LATIN)
        srfonts.set(qn("w:hAnsi"), LATIN)
        srfonts.set(qn("w:eastAsia"), CJK)

    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    # --- 封面 ---
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rk = kicker.add_run("PRODUCT & ARCHITECTURE PLAN")
    rk.bold = True
    rk.font.size = Pt(12)
    rk.font.color.rgb = ACCENT_H2
    set_cjk_font(rk)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    rt = title.add_run("通用 Agent 平台、AI Studio\n與 Manifund 專案研究整合")
    rt.bold = True
    rt.font.size = Pt(26)
    rt.font.color.rgb = ACCENT_H1
    set_cjk_font(rt)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(6)
    rs = sub.add_run("從金融垂直產品升級為可組合、可治理、可驗證的通用 AI 平台")
    rs.font.size = Pt(12)
    rs.font.color.rgb = GREY
    set_cjk_font(rs)

    meta_spacer = doc.add_paragraph()
    meta_spacer.paragraph_format.space_before = Pt(24)

    # 中繼資料表（取 markdown 首個 table）
    meta = next((b for b in blocks if b["type"] == "table"), None)
    if meta:
        rows = meta["rows"]
        t = doc.add_table(rows=len(rows), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            cells = row + [""] * (2 - len(row))
            for ci in range(2):
                cell = t.rows[ri].cells[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                add_inline(p, cells[ci], base_size=10)
                if ri == 0:
                    shade_cell(cell, "1F3A5F")
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for _ in range(2):
        doc.add_paragraph()
    prep = doc.add_paragraph()
    prep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp = prep.add_run("Prepared for product, architecture, security and UX review")
    rp.italic = True
    rp.font.size = Pt(9.5)
    rp.font.color.rgb = GREY
    set_cjk_font(rp)

    # --- 目錄（page_break_before 掛在內容段落上；獨立分頁空段落會在目錄滿版時形成整頁空白）---
    th = doc.add_paragraph()
    th.paragraph_format.page_break_before = True
    rth = th.add_run("目錄")
    rth.bold = True
    rth.font.size = Pt(16)
    rth.font.color.rgb = ACCENT_H1
    set_cjk_font(rth)
    add_toc(doc)

    # --- 正文（從第一個 ## 章節開始；# 標題與中繼資料表已於封面呈現）---
    body_start = next((idx for idx, b in enumerate(blocks) if b["type"] == "h2"), 0)
    code_block_seen = 0
    first_body_para = True

    for b in blocks[body_start:]:
        bt = b["type"]
        if bt == "h1":
            continue  # 文件標題已於封面呈現，正文不重複
        elif bt == "h2":
            p = doc.add_heading(level=1)
            if first_body_para:
                p.paragraph_format.page_break_before = True
                first_body_para = False
            add_inline(p, b["text"], base_size=18)
            for r in p.runs:
                r.font.color.rgb = ACCENT_H1
        elif bt == "h3":
            p = doc.add_heading(level=2)
            add_inline(p, b["text"], base_size=14)
            for r in p.runs:
                r.font.color.rgb = ACCENT_H2
        elif bt == "h4":
            p = doc.add_heading(level=3)
            add_inline(p, b["text"], base_size=12)
            for r in p.runs:
                r.font.color.rgb = ACCENT_H3
        elif bt == "para":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.first_line_indent = Pt(21)
            add_inline(p, b["text"])
        elif bt == "bullets":
            for lvl, txt in b["items"]:
                txt2 = txt
                style = "List Bullet" if lvl == 0 else "List Bullet 2"
                if re.match(r"\[[ xX]\]\s*", txt2):
                    box = "☑" if txt2[1] in "xX" else "☐"
                    txt2 = box + " " + re.sub(r"^\[[ xX]\]\s*", "", txt2)
                p = doc.add_paragraph(style=style)
                p.paragraph_format.line_spacing = 1.3
                add_inline(p, txt2)
        elif bt == "numbered":
            # 不用 "List Number" 樣式：python-docx 所有清單共用同一編號定義，
            # 會跨清單連續編號（第二份清單變 9–13 而非 1–5）。
            # 直接輸出 markdown 原始編號 + 懸掛縮排，保證每份清單從自己的號碼開始。
            for num, txt in b["items"]:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.75)
                add_inline(p, f"{num}. {txt}")
        elif bt == "quote":
            p = doc.add_paragraph(style="Intense Quote")
            p.paragraph_format.line_spacing = 1.3
            add_inline(p, b["text"])
        elif bt == "table":
            rows = b["rows"]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.rows[ri].cells[ci]
                    val = row[ci] if ci < len(row) else ""
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.2
                    add_inline(p, val, base_size=9.5)
                    if ri == 0:
                        shade_cell(cell, "E8EEF5")
                        for r in p.runs:
                            r.bold = True
            doc.add_paragraph()
        elif bt == "code":
            lang = b.get("lang", "")
            if lang == "mermaid":
                code_block_seen += 1
                if code_block_seen == 1 and ARCH_IMG.exists():
                    picp = doc.add_paragraph()
                    picp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = picp.add_run()
                    run.add_picture(str(ARCH_IMG), width=Cm(16))
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rc = cap.add_run("（架構流程圖）")
                    rc.italic = True
                    rc.font.size = Pt(9)
                    rc.font.color.rgb = GREY
                    set_cjk_font(rc)
                else:
                    cap = doc.add_paragraph()
                    rc = cap.add_run("（流程圖原始碼，以等寬呈現）")
                    rc.italic = True
                    rc.font.size = Pt(9)
                    rc.font.color.rgb = GREY
                    set_cjk_font(rc)
                    for ln in b["lines"]:
                        _add_code_line(doc, ln)
            else:
                for ln in b["lines"]:
                    _add_code_line(doc, ln)
        elif bt == "hr":
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "C0C0C0")
            pbdr.append(bottom)
            ppr.append(pbdr)

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DST)
    print(f"已產生：{DST}")


def _add_code_line(doc, line: str):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.first_line_indent = Pt(0)
    para_shade(p, "F4F5F7")
    r = p.add_run(line if line else " ")
    set_cjk_font(r, mono=True)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x24, 0x29, 0x33)


if __name__ == "__main__":
    if not SRC.exists():
        sys.exit(f"找不到來源：{SRC}")
    build()
